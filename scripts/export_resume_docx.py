#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 Markdown 文档导出为 Word（.docx）。

用途：简历类文档需要交 Word/PDF 时，直接把 docs/*.md 转成可编辑的 docx，
避免手工复制导致格式与数字走样。

支持：`#`~`###` 标题、`>` 引用、`-` 无序列表、`1.` 有序列表、`**粗体**`、
`` `行内代码` ``、`|` 表格、``` 代码块、`[文字](链接)`、`<!-- -->` 注释（跳过）、`---` 分隔线。

用法：
    python scripts/export_resume_docx.py                      # 默认导出简历
    python scripts/export_resume_docx.py 输入.md 输出.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_SRC = BASE_DIR / "docs" / "简历项目经验-STAR.md"
DEFAULT_DST = BASE_DIR / "docs" / "简历项目经验-STAR.docx"

COMMENT_RE = re.compile(r"<!--.*?-->", re.S)
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
ORDERED_RE = re.compile(r"^\s*(\d+)\.\s+(.*)$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _add_runs(paragraph, text: str) -> None:
    """按 Markdown 行内标记写入 run：**粗体**、`代码`。"""
    text = LINK_RE.sub(lambda m: f"{m.group(1)}（{m.group(2)}）", text)
    pos = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        bold, code = match.group(1), match.group(2)
        if bold is not None:
            run = paragraph.add_run(bold)
            run.bold = True
        else:
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _split_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def _is_separator(line: str) -> bool:
    return bool(re.match(r"^\|[\s:|-]+\|$", line.strip()))


def markdown_to_docx(src: Path, dst: Path) -> dict:
    raw = src.read_text(encoding="utf-8")
    raw = COMMENT_RE.sub("", raw)
    lines = raw.splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)

    stats = {"headings": 0, "paragraphs": 0, "bullets": 0, "tables": 0, "code_blocks": 0}

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for code_line in block:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                run = para.add_run(code_line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            stats["code_blocks"] += 1
            continue

        if not stripped:
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped):
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2).strip(), level=level)
            stats["headings"] += 1
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and _is_separator(lines[i + 1]):
            header = _split_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for idx, text in enumerate(header):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                _add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    value = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    _add_runs(cells[idx].paragraphs[0], value)
            stats["tables"] += 1
            doc.add_paragraph()
            continue

        if stripped.startswith(">"):
            quote: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for text in quote:
                if not text:
                    continue
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Pt(18)
                para.paragraph_format.space_after = Pt(2)
                _add_runs(para, text)
                for run in para.runs:
                    run.italic = True
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            stats["paragraphs"] += 1
            continue

        bullet = BULLET_RE.match(line)
        ordered = ORDERED_RE.match(line)
        if bullet or ordered:
            text = bullet.group(1) if bullet else ordered.group(2)
            style_name = "List Bullet" if bullet else "List Number"
            para = doc.add_paragraph(style=style_name)
            _add_runs(para, text)
            stats["bullets"] += 1
            i += 1
            continue

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_runs(para, stripped)
        stats["paragraphs"] += 1
        i += 1

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    stats["bytes"] = dst.stat().st_size
    return stats


def main(argv: list[str]) -> int:
    src = Path(argv[1]) if len(argv) > 1 else DEFAULT_SRC
    dst = Path(argv[2]) if len(argv) > 2 else DEFAULT_DST
    if not src.exists():
        print(f"[x] 源文件不存在: {src}")
        return 1
    stats = markdown_to_docx(src, dst)
    print(f"[ok] {src.name} -> {dst.name}")
    print(
        "     标题 {headings} / 段落 {paragraphs} / 列表 {bullets} / 表格 {tables} / 代码块 {code_blocks} / {bytes}B".format(
            **stats
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
